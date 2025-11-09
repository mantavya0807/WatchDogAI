"""
MCP Server for DOCX Processing
Extracts text from DOCX files and checks for PII using existing detection stack.
"""

import sys
from pathlib import Path

# Add parent directory to path for imports
sys.path.insert(0, str(Path(__file__).parent.parent))

from typing import Dict, Optional
import time

try:
    from fastmcp import FastMCP
except ImportError:
    print("Error: fastmcp not installed. Run: pip install fastmcp")
    sys.exit(1)

try:
    from docx import Document
except ImportError:
    print("Error: python-docx not installed. Run: pip install python-docx")
    sys.exit(1)

# Import existing obfuscator
from src.obfuscator import PIIObfuscator

# Initialize MCP server
mcp = FastMCP("pii-guard-docx")

# Initialize obfuscator (lazy loading)
_obfuscator: Optional[PIIObfuscator] = None


def get_obfuscator() -> PIIObfuscator:
    """Get or create obfuscator instance (lazy loading)"""
    global _obfuscator
    if _obfuscator is None:
        print("Initializing PII Obfuscator for DOCX processing...")
        _obfuscator = PIIObfuscator(
            use_regex=True,
            use_spacy=True,
            use_transformer=True,
            escrow_db_path="data/escrow/pii_escrow.db"
        )
    return _obfuscator


def check_docx(filepath: str, max_file_size: int = 50 * 1024 * 1024) -> Dict:
    """
    Extract text from DOCX file and check for PII.
    
    Args:
        filepath: Path to DOCX file
        max_file_size: Maximum file size in bytes (default 50MB)
        
    Returns:
        Dictionary with:
        - safe_text: Obfuscated text with PII replaced
        - original_text: Original extracted text
        - pii_found: Number of PII items detected
        - detections: List of detected PII types
        - extraction_time_ms: Time taken to extract text
        - detection_time_ms: Time taken to detect PII
        - obfuscation_time_ms: Time taken to obfuscate
    """
    start_time = time.time()
    
    try:
        # Check file size
        file_size = Path(filepath).stat().st_size
        if file_size > max_file_size:
            return {
                "error": f"File too large: {file_size / (1024*1024):.1f}MB (max {max_file_size / (1024*1024):.1f}MB)",
                "pii_found": 0
            }
        
        # Check if file exists
        if not Path(filepath).exists():
            return {
                "error": f"File not found: {filepath}",
                "pii_found": 0
            }
        # Step 1: Extract text from DOCX and process each paragraph individually
        print(f"Extracting text from DOCX: {filepath}")
        doc = Document(filepath)
        
        # Store paragraphs and tables with their original text for mapping
        paragraph_texts = []
        table_row_texts = []
        
        # Extract paragraphs
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                paragraph_texts.append(paragraph.text)
        
        # Extract tables
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    if cell.text.strip():
                        row_text.append(cell.text.strip())
                if row_text:
                    table_row_texts.append(" | ".join(row_text))
        
        all_text_parts = paragraph_texts + table_row_texts
        full_text = "\n\n".join(all_text_parts)
        extraction_time = (time.time() - start_time) * 1000
        
        if not full_text.strip():
            return {
                "safe_text": "",
                "original_text": "",
                "pii_found": 0,
                "detections": [],
                "extraction_time_ms": extraction_time,
                "detection_time_ms": 0.0,
                "obfuscation_time_ms": 0.0,
                "error": "No text found in DOCX file"
            }
        
        print(f"Extracted {len(full_text)} characters from DOCX")
        
        # Step 2: Detect and obfuscate PII using existing detection stack
        # The obfuscator follows: transformer → spaCy → regex
        print("Detecting PII in extracted text...")
        print("Using obfuscator flow: Transformer → spaCy → Regex")
        obfuscator = get_obfuscator()
        
        # Process each text part individually to maintain mapping
        obfuscated_parts = []
        all_replacements = {}
        detection_types = set()
        total_detection_time = 0.0
        total_obfuscation_time = 0.0
        
        for text_part in all_text_parts:
            # Obfuscate each part individually
            result = obfuscator.obfuscate(text_part, source=f"docx:{filepath}")
            obfuscated_parts.append(result.obfuscated_text)
            
            # Collect replacements and detection types
            for placeholder, original in result.replacements.items():
                all_replacements[placeholder] = original
                if placeholder.startswith("{") and placeholder.endswith("}"):
                    entity_type = placeholder[1:].split("_")[0]
                    detection_types.add(entity_type)
            
            total_detection_time += result.detection_time_ms
            total_obfuscation_time += result.obfuscation_time_ms
        
        # Combine obfuscated text
        obfuscated_full_text = "\n\n".join(obfuscated_parts)
        
        # Step 3: Create new DOCX document with obfuscated text
        # Map obfuscated text back to document structure
        text_to_obfuscated = dict(zip(all_text_parts, obfuscated_parts))
        
        # Create new document
        safe_doc = Document()
        
        # Reconstruct paragraphs with obfuscated text
        para_idx = 0
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                if para_idx < len(paragraph_texts):
                    original_para = paragraph_texts[para_idx]
                    obfuscated_para = text_to_obfuscated.get(original_para, original_para)
                    safe_doc.add_paragraph(obfuscated_para)
                    para_idx += 1
                else:
                    safe_doc.add_paragraph(paragraph.text)
        
        # Reconstruct tables with obfuscated text
        table_row_idx = 0
        for table in doc.tables:
            safe_table = safe_doc.add_table(rows=len(table.rows), cols=len(table.columns))
            for row_idx, row in enumerate(table.rows):
                row_text_parts = []
                for col_idx, cell in enumerate(row.cells):
                    if cell.text.strip():
                        row_text_parts.append(cell.text.strip())
                
                if row_text_parts:
                    # Join row cells with " | " to match extraction format
                    row_text = " | ".join(row_text_parts)
                    if table_row_idx < len(table_row_texts):
                        original_row = table_row_texts[table_row_idx]
                        obfuscated_row = text_to_obfuscated.get(original_row, original_row)
                        table_row_idx += 1
                    else:
                        obfuscated_row = row_text
                    
                    # Split obfuscated row back into cells
                    if " | " in obfuscated_row:
                        cell_values = obfuscated_row.split(" | ")
                        for col_idx, cell_value in enumerate(cell_values):
                            if col_idx < len(safe_table.rows[row_idx].cells):
                                safe_table.rows[row_idx].cells[col_idx].text = cell_value
                    else:
                        # If no separator, put all text in first cell
                        if len(safe_table.rows[row_idx].cells) > 0:
                            safe_table.rows[row_idx].cells[0].text = obfuscated_row
        
        total_time = (time.time() - start_time) * 1000
        
        return {
            "safe_text": obfuscated_full_text,
            "safe_docx": safe_doc,  # Document object for saving
            "original_text": full_text,
            "pii_found": len(all_replacements),
            "detections": sorted(list(detection_types)),
            "extraction_time_ms": extraction_time,
            "detection_time_ms": total_detection_time,
            "obfuscation_time_ms": total_obfuscation_time,
            "total_time_ms": total_time
        }
        
    except FileNotFoundError:
        return {
            "error": f"File not found: {filepath}",
            "pii_found": 0
        }
    except Exception as e:
        return {
            "error": f"Error processing DOCX: {str(e)}",
            "pii_found": 0
        }


# Register as MCP tool
mcp.tool()(check_docx)


def extract_docx_text(filepath: str) -> Dict:
    """
    Extract text from DOCX without PII detection (for testing/debugging).
    
    Args:
        filepath: Path to DOCX file
        
    Returns:
        Dictionary with extracted text and metadata
    """
    try:
        doc = Document(filepath)
        
        text_parts = []
        
        # Extract paragraphs
        for para_num, paragraph in enumerate(doc.paragraphs, 1):
            if paragraph.text.strip():
                text_parts.append(f"--- Paragraph {para_num} ---\n{paragraph.text}")
        
        # Extract tables
        for table_num, table in enumerate(doc.tables, 1):
            table_text = f"--- Table {table_num} ---\n"
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    if cell.text.strip():
                        row_text.append(cell.text.strip())
                if row_text:
                    table_text += " | ".join(row_text) + "\n"
            if table_text.strip():
                text_parts.append(table_text)
        
        full_text = "\n\n".join(text_parts)
        
        return {
            "text": full_text,
            "char_count": len(full_text),
            "paragraph_count": len([p for p in doc.paragraphs if p.text.strip()]),
            "table_count": len(doc.tables),
            "extracted": True
        }
        
    except Exception as e:
        return {
            "error": f"Error extracting text: {str(e)}",
            "extracted": False
        }


# Register as MCP tool
mcp.tool()(extract_docx_text)


if __name__ == "__main__":
    # Run MCP server via stdio
    print("=" * 70)
    print("PII Guard - DOCX MCP Server")
    print("=" * 70)
    print("Starting server via stdio transport...")
    print("This server extracts text from DOCX files and checks for PII")
    print("=" * 70)
    mcp.run()

