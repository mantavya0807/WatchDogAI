"""
Integration Tests for PII Guard MCP Servers
Tests unified server, file type detection, CLI options, and error handling.
"""

import sys
import os
import io
from pathlib import Path
import subprocess
import tempfile
import shutil

# Fix Unicode encoding for Windows console
if sys.platform == 'win32':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')

# Add parent directory to path
sys.path.insert(0, str(Path(__file__).parent))

from mcp_servers.unified_server import check_file, extract_text, list_supported_types, detect_file_type
from mcp_servers.code_server import check_code
from mcp_servers.docx_server import check_docx
from mcp_servers.pdf_server import check_pdf

# Test results
test_results = {
    "passed": [],
    "failed": [],
    "warnings": []
}

def log_test(name, passed, message=""):
    """Log test result"""
    if passed:
        test_results["passed"].append(name)
        print(f"✅ PASS: {name}")
        if message:
            print(f"   {message}")
    else:
        test_results["failed"].append(name)
        print(f"❌ FAIL: {name}")
        if message:
            print(f"   {message}")

def log_warning(name, message):
    """Log warning"""
    test_results["warnings"].append(f"{name}: {message}")
    print(f"⚠️  WARN: {name} - {message}")

print("=" * 70)
print("PII Guard MCP Servers - Integration Tests")
print("=" * 70)
print()

# Test 1: List supported file types
print("\n" + "=" * 70)
print("TEST 1: List Supported File Types")
print("=" * 70)
try:
    result = list_supported_types()
    has_pdf = 'pdf' in str(result.get('supported_types', {}))
    has_code = 'code' in str(result.get('supported_types', {}))
    has_docx = 'docx' in str(result.get('supported_types', {}))
    
    log_test("List supported types", has_pdf and has_code and has_docx,
             f"Found: {result.get('total_extensions', 0)} extensions")
except Exception as e:
    log_test("List supported types", False, str(e))

# Test 2: File type detection
print("\n" + "=" * 70)
print("TEST 2: File Type Detection")
print("=" * 70)

test_files = {
    "test.py": "code",
    "test.js": "code",
    "test.java": "code",
    "test.pdf": "pdf",
    "test.docx": "docx",
    "test.txt": None,  # Not supported
    "test.unknown": None  # Not supported
}

for filename, expected_type in test_files.items():
    detected = detect_file_type(filename)
    passed = detected == expected_type
    log_test(f"Detect {filename}", passed,
             f"Expected: {expected_type}, Got: {detected}")

# Test 3: Create test files
print("\n" + "=" * 70)
print("TEST 3: Creating Test Files")
print("=" * 70)

test_dir = Path("test_integration_files")
test_dir.mkdir(exist_ok=True)

# Create test code file
test_code_file = test_dir / "test_code.py"
test_code_file.write_text('''"""
Test Code File with PII
"""
API_KEY = "sk-1234567890abcdef"
USER_EMAIL = "john.smith@example.com"
PHONE_NUMBER = "555-123-4567"
# Comment with SSN: 123-45-6789
''')
log_test("Create test code file", test_code_file.exists())

# Create test DOCX file
try:
    from docx import Document
    test_docx_file = test_dir / "test_docx.docx"
    doc = Document()
    doc.add_paragraph("Test Document with PII")
    doc.add_paragraph("Contact John Smith at john.smith@example.com")
    doc.add_paragraph("Phone: (555) 123-4567")
    doc.add_paragraph("SSN: 123-45-6789")
    doc.save(str(test_docx_file))
    log_test("Create test DOCX file", test_docx_file.exists())
except Exception as e:
    log_test("Create test DOCX file", False, str(e))

# Create test PDF file
try:
    import fitz
    test_pdf_file = test_dir / "test_pdf.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "Test PDF with PII\nContact John Smith at john.smith@example.com\nPhone: (555) 123-4567\nSSN: 123-45-6789", fontsize=11)
    doc.save(str(test_pdf_file))
    doc.close()
    log_test("Create test PDF file", test_pdf_file.exists())
except Exception as e:
    log_test("Create test PDF file", False, str(e))

# Test 4: Unified server - Code files
print("\n" + "=" * 70)
print("TEST 4: Unified Server - Code Files")
print("=" * 70)

if test_code_file.exists():
    try:
        result = check_file(str(test_code_file), use_tree_sitter=True)
        has_pii = result.get("pii_found", 0) > 0
        has_safe_text = "safe_text" in result
        has_detections = "detections" in result
        
        log_test("Check code file", not result.get("error") and has_pii and has_safe_text,
                 f"PII found: {result.get('pii_found', 0)}, Types: {result.get('detections', [])}")
        
        # Verify output format
        if has_safe_text:
            safe_text = result["safe_text"]
            has_placeholders = "{" in safe_text and "}" in safe_text
            log_test("Code output has placeholders", has_placeholders,
                     f"Sample: {safe_text[:100]}...")
    except Exception as e:
        log_test("Check code file", False, str(e))
else:
    log_warning("Check code file", "Test file not created")

# Test 5: Unified server - DOCX files
print("\n" + "=" * 70)
print("TEST 5: Unified Server - DOCX Files")
print("=" * 70)

test_docx_file = test_dir / "test_docx.docx"
if test_docx_file.exists():
    try:
        result = check_file(str(test_docx_file))
        has_pii = result.get("pii_found", 0) > 0
        has_safe_text = "safe_text" in result
        has_safe_docx = "safe_docx" in result
        
        log_test("Check DOCX file", not result.get("error") and has_pii and has_safe_text,
                 f"PII found: {result.get('pii_found', 0)}, Types: {result.get('detections', [])}")
        
        log_test("DOCX output has document object", has_safe_docx,
                 "Can save as DOCX file")
    except Exception as e:
        log_test("Check DOCX file", False, str(e))
else:
    log_warning("Check DOCX file", "Test file not created")

# Test 6: Unified server - PDF files
print("\n" + "=" * 70)
print("TEST 6: Unified Server - PDF Files")
print("=" * 70)

test_pdf_file = test_dir / "test_pdf.pdf"
if test_pdf_file.exists():
    try:
        result = check_file(str(test_pdf_file))
        has_pii = result.get("pii_found", 0) > 0
        has_safe_text = "safe_text" in result
        has_safe_pdf = "safe_pdf" in result
        
        log_test("Check PDF file", not result.get("error") and has_pii and has_safe_text,
                 f"PII found: {result.get('pii_found', 0)}, Types: {result.get('detections', [])}")
        
        log_test("PDF output has document object", has_safe_pdf,
                 "Can save as PDF file")
    except Exception as e:
        log_test("Check PDF file", False, str(e))
else:
    log_warning("Check PDF file", "Test file not created")

# Test 7: Error handling - File not found
print("\n" + "=" * 70)
print("TEST 7: Error Handling - File Not Found")
print("=" * 70)

try:
    result = check_file("nonexistent_file.py")
    has_error = "error" in result
    error_message = result.get("error", "")
    log_test("Handle file not found", has_error and "not found" in error_message.lower(),
             f"Error: {error_message}")
except Exception as e:
    log_test("Handle file not found", False, str(e))

# Test 8: Error handling - Unsupported file type
print("\n" + "=" * 70)
print("TEST 8: Error Handling - Unsupported File Type")
print("=" * 70)

unsupported_file = test_dir / "test.unknown"
unsupported_file.write_text("test content")
try:
    result = check_file(str(unsupported_file))
    has_error = "error" in result
    error_message = result.get("error", "")
    log_test("Handle unsupported file type", has_error and ("unsupported" in error_message.lower() or "not supported" in error_message.lower()),
             f"Error: {error_message}")
except Exception as e:
    log_test("Handle unsupported file type", False, str(e))

# Test 9: Error handling - Empty file
print("\n" + "=" * 70)
print("TEST 9: Error Handling - Empty File")
print("=" * 70)

empty_file = test_dir / "empty.py"
empty_file.write_text("")
try:
    result = check_file(str(empty_file))
    # Empty file might return no PII or an error - both are acceptable
    has_result = "pii_found" in result or "error" in result
    log_test("Handle empty file", has_result,
             f"Result: {result.get('pii_found', 0)} PII found or error: {result.get('error', 'None')}")
except Exception as e:
    log_test("Handle empty file", False, str(e))

# Test 10: CLI client - Check command
print("\n" + "=" * 70)
print("TEST 10: CLI Client - Check Command")
print("=" * 70)

if test_code_file.exists():
    try:
        result = subprocess.run(
            [sys.executable, "mcp_client.py", "check", str(test_code_file), "-o", str(test_dir / "output_code.py")],
            capture_output=True,
            text=True,
            encoding='utf-8',
            errors='replace',
            timeout=60
        )
        success = result.returncode == 0 and (result.stdout and ("PII DETECTED" in result.stdout or "PII" in result.stdout))
        output_file = test_dir / "output_code.py"
        has_output = output_file.exists()
        
        log_test("CLI check command", success and has_output,
                 f"Exit code: {result.returncode}, Output file exists: {has_output}")
        
        if has_output:
            # Verify output is Python file
            content = output_file.read_text()
            has_placeholders = "{" in content and "}" in content
            log_test("CLI output is correct format", has_placeholders,
                     "Output contains placeholders")
    except subprocess.TimeoutExpired:
        log_test("CLI check command", False, "Command timed out")
    except Exception as e:
        log_test("CLI check command", False, str(e))

# Test 11: CLI client - List command
print("\n" + "=" * 70)
print("TEST 11: CLI Client - List Command")
print("=" * 70)

try:
    result = subprocess.run(
        [sys.executable, "mcp_client.py", "list"],
        capture_output=True,
        text=True,
        encoding='utf-8',
        errors='replace',
        timeout=30
    )
    success = result.returncode == 0 and (result.stdout and "SUPPORTED" in result.stdout)
    log_test("CLI list command", success,
             f"Exit code: {result.returncode}")
except subprocess.TimeoutExpired:
    log_test("CLI list command", False, "Command timed out")
except Exception as e:
    log_test("CLI list command", False, str(e))

# Test 12: CLI client - Verbose option
print("\n" + "=" * 70)
print("TEST 12: CLI Client - Verbose Option")
print("=" * 70)

if test_code_file.exists():
    try:
        result = subprocess.run(
            [sys.executable, "mcp_client.py", "check", str(test_code_file), "-v"],
            capture_output=True,
            text=True,
            encoding='utf-8',
            errors='replace',
            timeout=60
        )
        success = result.returncode == 0 and (result.stdout and ("Original Text" in result.stdout or "Safe Text" in result.stdout or "PII" in result.stdout))
        log_test("CLI verbose option", success,
                 f"Exit code: {result.returncode}")
    except subprocess.TimeoutExpired:
        log_test("CLI verbose option", False, "Command timed out")
    except Exception as e:
        log_test("CLI verbose option", False, str(e))

# Test 13: File extension preservation
print("\n" + "=" * 70)
print("TEST 13: File Extension Preservation")
print("=" * 70)

if test_code_file.exists():
    try:
        # Test with output file without extension
        output_no_ext = test_dir / "output_code"
        result = subprocess.run(
            [sys.executable, "mcp_client.py", "check", str(test_code_file), "-o", str(output_no_ext)],
            capture_output=True,
            text=True,
            encoding='utf-8',
            errors='replace',
            timeout=60
        )
        # Should create output_code.py (preserves .py extension)
        output_with_ext = test_dir / "output_code.py"
        has_correct_ext = output_with_ext.exists()
        
        log_test("Preserve file extension", has_correct_ext,
                 f"Output file: {output_with_ext.name if has_correct_ext else 'not found'}")
    except subprocess.TimeoutExpired:
        log_test("Preserve file extension", False, "Command timed out")
    except Exception as e:
        log_test("Preserve file extension", False, str(e))

# Test 14: Individual server functions
print("\n" + "=" * 70)
print("TEST 14: Individual Server Functions")
print("=" * 70)

if test_code_file.exists():
    try:
        result = check_code(str(test_code_file), use_tree_sitter=True)
        has_pii = result.get("pii_found", 0) > 0
        log_test("Code server function", has_pii,
                 f"PII found: {result.get('pii_found', 0)}")
    except Exception as e:
        log_test("Code server function", False, str(e))

if test_docx_file.exists():
    try:
        result = check_docx(str(test_docx_file))
        has_pii = result.get("pii_found", 0) > 0
        log_test("DOCX server function", has_pii,
                 f"PII found: {result.get('pii_found', 0)}")
    except Exception as e:
        log_test("DOCX server function", False, str(e))

if test_pdf_file.exists():
    try:
        result = check_pdf(str(test_pdf_file))
        has_pii = result.get("pii_found", 0) > 0
        log_test("PDF server function", has_pii,
                 f"PII found: {result.get('pii_found', 0)}")
    except Exception as e:
        log_test("PDF server function", False, str(e))

# Summary
print("\n" + "=" * 70)
print("TEST SUMMARY")
print("=" * 70)
print(f"✅ Passed: {len(test_results['passed'])}")
print(f"❌ Failed: {len(test_results['failed'])}")
print(f"⚠️  Warnings: {len(test_results['warnings'])}")
print()

if test_results['failed']:
    print("Failed Tests:")
    for test in test_results['failed']:
        print(f"  - {test}")
    print()

if test_results['warnings']:
    print("Warnings:")
    for warning in test_results['warnings']:
        print(f"  - {warning}")
    print()

# Cleanup
print("Cleaning up test files...")
if test_dir.exists():
    try:
        shutil.rmtree(test_dir)
        print("✅ Test files cleaned up")
    except Exception as e:
        print(f"⚠️  Could not clean up test files: {e}")

# Final result
if len(test_results['failed']) == 0:
    print("\n✅ All tests passed!")
    sys.exit(0)
else:
    print(f"\n❌ {len(test_results['failed'])} test(s) failed")
    sys.exit(1)

